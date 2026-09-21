#!/usr/bin/env python3
"""Verify the generated report: structure, figures, page breaks and leftovers."""

from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from docx import Document

PATH = Path(sys.argv[1] if len(sys.argv) > 1 else "Cryptographic_Cipher_Analyzer_Project_Report.docx")

LEFTOVERS = [
    "STUDENT INSTRUCTION",
    "INSERT SCREENSHOT",
    "Write your",
    "Objective 1 - e.g.",
    "Learning outcome 1",
    "[ Finding",
    "[ Challenge",
    "[ Tool ",
    "[ Resource",
    "[ Description",
    "[ How you",
    "[ Solution ]",
    "[ Add more",
    "[ e.g.",
    "[ Version ]",
    "[ Purpose ]",
]


def main() -> int:
    document = Document(PATH)
    paragraphs = document.paragraphs
    cell_text = "\n".join(
        paragraph.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    )
    text = "\n".join(p.text for p in paragraphs) + "\n" + cell_text

    print(f"File            : {PATH}")
    print(f"Size            : {PATH.stat().st_size / 1024 / 1024:.2f} MB")
    print(f"Paragraphs      : {len(paragraphs)}")
    print(f"Tables          : {len(document.tables)}")
    print(f"Inline images   : {len(document.inline_shapes)}")

    captions = re.findall(r"^Figure (\d+): ", text, re.M)
    print(f"Figure captions : {len(captions)} -> {captions}")
    missing = sorted(set(range(1, len(captions) + 1)) - {int(c) for c in captions})
    if missing:
        print(f"  !! missing figure numbers: {missing}")

    banners = [p.text for p in paragraphs if p.text.strip().startswith(("SECTION", "DECLARATION", "CONTENTS"))]
    print("Section banners :")
    for banner in banners:
        print(f"  - {banner}")

    steps = [p.text for p in paragraphs if re.match(r"^Step \d+ [\u2013-] ", p.text)]
    print(f"Methodology steps: {len(steps)}")
    for number, name in enumerate(steps, start=1):
        if not re.match(rf"^Step {number} [\u2013-] ", name):
            print(f"  !! out of order: {name!r}")

    xml = document.element.body.xml
    manual = len(re.findall(r'w:type="page"', xml))
    before = len(re.findall(r"<w:pageBreakBefore[^/]*/>", xml))
    print(f"Page breaks     : {before} (page-break-before) + {manual} (manual)")

    with zipfile.ZipFile(PATH) as archive:
        media = [n for n in archive.namelist() if n.startswith("word/media/")]
    print(f"Embedded media  : {len(media)}")

    hits = [(phrase, text.count(phrase)) for phrase in LEFTOVERS if phrase in text]
    print(f"Leftover template text: {hits if hits else 'none'}")

    placeholders = sorted(set(re.findall(r"\[[^\]\n]{3,60}\]", text)))
    print(f"Fill-in placeholders : {placeholders}")

    mismatched = re.findall(r"Figure (\d+): [^\n]{0,40}", text)
    print("\nAll good" if not missing else "\nProblems found")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
