#!/usr/bin/env python3
"""Verify the generated presentation: slides, images, notes, dimensions."""

from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "Cryptographic_Cipher_Analyzer_Presentation.pptx"

prs = Presentation(str(PATH))
SW, SH = prs.slide_width, prs.slide_height
print(f"Slide size : {SW / 914400:.2f} x {SH / 914400:.2f} in")
print(f"Slides     : {len(prs.slides)}")
print("-" * 72)

total_pics = 0
problems = []
for idx, slide in enumerate(prs.slides, start=1):
    pics = sum(1 for sh in slide.shapes if sh.shape_type == 13)  # PICTURE
    total_pics += pics

    # notes
    notes = ""
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text.strip()
    n_words = len(notes.split())

    # overflow check: any shape extending beyond slide bounds
    oob = []
    for sh in slide.shapes:
        try:
            if sh.left is None:
                continue
            if sh.left < -9525 or sh.top < -9525 or \
               sh.left + sh.width > SW + 9525 or sh.top + sh.height > SH + 9525:
                oob.append(sh.shape_type)
        except (TypeError, AttributeError):
            continue

    title = ""
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            txt = sh.text_frame.text.strip().splitlines()[0]
            if len(txt) > 8 and not txt.startswith("$"):
                title = txt[:52]
                break

    flag = ""
    if oob:
        flag += f"  !! {len(oob)} out-of-bounds"
        problems.append((idx, oob))
    if not notes:
        flag += "  !! no notes"
        problems.append((idx, "missing notes"))

    print(f"{idx:02d} | pics={pics} | notes={n_words:4d}w | {title}{flag}")

print("-" * 72)
print(f"Total embedded screenshots: {total_pics}")
if problems:
    print(f"PROBLEMS on slides: {[p[0] for p in problems]}")
    raise SystemExit(1)
print("ALL CHECKS PASSED")
